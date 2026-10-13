"""Master resume: extraction, grouping, questions, verification, confirmation, rendering."""

from pathlib import Path

import docx
import pytest

from applier.config import ConfigError
from applier.profile import load_profile
from applier.resume.build import build_master
from applier.resume.extract import extract
from applier.resume.model import GroupsFile, load_master, require_confirmed, save_master
from applier.resume.render import RenderError, base_document, render_pdf, resume_filename
from applier.resume.sources import read_source
from applier.resume.verify import style_problems, verify_master
from tests.conftest import FIXTURES

FAKE_A = [
    ("Test Person", None),
    ("test@example.com | +1 555 010 0000", None),
    ("EDUCATION", None),
    ("State University, Master of Things\tBoston, MA | 2020 - 2022", None),
    ("GPA 4.0/4.0; Dean's List", "List Bullet"),
    ("EXPERIENCE", None),
    ("Acme Labs (analytics startup)", None),
    ("Analyst\tNew York, NY | 2023 - Present", None),
    ("Built a forecasting model covering 40 stores that cut stockouts by 12%", "List Bullet"),
    ("Met the President of Acme to present findings", "List Bullet"),
    ("SKILLS", None),
    ("Technical: R (advanced); Python (intermediate)", None),
]
FAKE_B = [
    ("Test Person", None),
    ("test@example.com", None),
    ("EXPERIENCE", None),
    ("Acme Labs", None),
    ("Analyst\tRemote | 2023 - Present", None),
    ("Built a forecasting model covering 40 stores that cut stockouts by 12%", "List Bullet"),
    ("Built a demand model for 40 stores in Excel", "List Bullet"),
    ("Met the President's office of Acme to present findings", "List Bullet"),
    ("Note to recruiter: available immediately", "List Bullet"),
]


def make_docx(path: Path, lines) -> Path:
    document = docx.Document()
    for text, style in lines:
        document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
    document.save(str(path))
    return path


@pytest.fixture
def sources(tmp_path):
    folder = tmp_path / "source"
    folder.mkdir()
    make_docx(folder / "a.docx", FAKE_A)
    make_docx(folder / "b.docx", FAKE_B)
    return folder


def spec(**overrides) -> GroupsFile:
    data = {
        "name": "Test Person",
        "location": "Boston, MA",
        "education": [{
            "id": "state", "school": "State University", "location": "Boston, MA", "dates": "2020 - 2022",
            "groups": [{"id": "state_honors", "label": "Honors", "rules": [{"context": ["State University"]}]}],
        }],
        "experience": [{
            "id": "acme", "organization": "Acme Labs", "title": "Analyst", "location": "New York, NY",
            "dates": "2023 - Present",
            "groups": [
                {"id": "acme_model", "label": "Forecasting model", "rules": [{"match": ["forecasting model", "demand model"]}]},
                {"id": "acme_other", "label": "Other", "rules": [{"context": ["Acme"]}]},
            ],
        }],
        "skills": [{"id": "sk_tech", "label": "Technical", "rules": [{"match": ["^Technical:"]}]}],
        "ignore": [{"match": ["Note to recruiter"]}],
        "conflicts": [{
            "id": "c_president",
            "question": "The President, or the President's office?",
            "groups": ["acme_other"],
            "options": {
                "a": {"label": "The office", "pattern": "President's office"},
                "b": {"label": "The President", "pattern": "President"},
            },
        }],
    }
    data.update(overrides)
    return GroupsFile.model_validate(data)


def built(sources, groups=None):
    docs = [read_source(p) for p in sorted(sources.iterdir())]
    master, report = build_master(groups or spec(), extract(docs), [d.name for d in docs])
    return master, report


def variants(master, group_id):
    return {v.text: v for g in master.all_groups() if g.id == group_id for v in g.variants}


def test_word_files_are_read_into_bullets_and_headings(sources):
    doc = read_source(sources / "a.docx")
    kinds = [(p.kind, p.text[:20]) for p in doc.paragraphs]
    assert kinds[0] == ("name", "Test Person")
    assert kinds[1][0] == "contact"
    bullet = next(p for p in doc.paragraphs if p.text.startswith("Built a forecasting"))
    assert bullet.kind == "bullet"
    assert bullet.context == "Acme Labs (analytics startup) / Analyst"
    assert any(p.kind == "text" and p.text.startswith("Technical:") for p in doc.paragraphs)


def test_identical_wordings_are_listed_once_with_every_file(sources):
    docs = [read_source(p) for p in sorted(sources.iterdir())]
    shared = [p for p in extract(docs) if p.text.startswith("Built a forecasting model")]
    assert len(shared) == 1
    assert shared[0].sources == ["a.docx", "b.docx"]


def test_wordings_are_grouped_and_nothing_is_lost(sources):
    master, report = built(sources)
    assert not report.unassigned
    assert report.ignored == 1  # the note to a recruiter
    assert set(variants(master, "acme_model")) == {
        "Built a forecasting model covering 40 stores that cut stockouts by 12%",
        "Built a demand model for 40 stores in Excel",
    }
    assert "Technical: R (advanced); Python (intermediate)" in variants(master, "sk_tech")
    assert "GPA 4.0/4.0; Dean's List" in variants(master, "state_honors")


def test_unanswered_question_keeps_wordings_out_of_use(sources):
    master, report = built(sources)
    others = variants(master, "acme_other")
    assert {v.status for v in others.values()} == {"disputed"}
    assert report.disputed == {"c_president": 2}


def test_answer_keeps_only_consistent_wordings(sources):
    groups = spec()
    groups.conflicts[0].resolution = "a"  # the office, not the President
    master, _ = built(sources, groups)
    others = variants(master, "acme_other")
    assert others["Met the President's office of Acme to present findings"].status == "ok"
    assert others["Met the President of Acme to present findings"].status == "excluded"


def test_wordings_can_be_excluded_by_id(sources):
    master, _ = built(sources)
    target = variants(master, "acme_model")["Built a demand model for 40 stores in Excel"]
    master, _ = built(sources, spec(exclude_ids={target.id: "user said Excel was not used"}))
    assert variants(master, "acme_model")[target.text].status == "excluded"


def test_facts_stated_in_chat_join_their_group(sources):
    groups = spec(added=[{"group": "acme_model", "stated": "2026-10-12",
                          "text": "Extended the model to 12 new stores in 2026"}])
    groups.conflicts[0].resolution = "a"
    master, _ = built(sources, groups)
    added = variants(master, "acme_model")["Extended the model to 12 new stores in 2026"]
    assert added.sources == ["stated by you on 2026-10-12"] and added.status == "ok"
    assert verify_master(master, sources).ok  # no resume file needed for a stated fact

    with pytest.raises(ValueError, match="unknown group"):
        built(sources, spec(added=[{"group": "nope", "stated": "2026-10-12", "text": "x"}]))


def test_summary_lines_are_dropped_when_not_wanted(tmp_path):
    folder = tmp_path / "src"
    folder.mkdir()
    make_docx(folder / "s.docx", [("Test Person", None), ("test@example.com", None),
                                  ("Analyst with six years of experience in everything.", None), *FAKE_A[2:]])
    master, report = built(folder)
    assert master.summaries is None
    assert not report.unassigned
    assert "six years" not in " ".join(v.text for g in master.all_groups() for v in g.variants)


def test_unmatched_paragraphs_are_reported(sources):
    groups = spec()
    groups.experience[0].groups = groups.experience[0].groups[:1]  # drop the catch-all
    _, report = built(sources, groups)
    assert {p.text for p in report.unassigned} == {
        "Met the President of Acme to present findings",
        "Met the President's office of Acme to present findings",
    }


def test_verify_catches_a_planted_claim(sources):
    """Every wording must appear word for word in the files it lists."""
    groups = spec()
    groups.conflicts[0].resolution = "a"
    master, _ = built(sources, groups)
    assert verify_master(master, sources).ok
    planted = master.experience[0].groups[0].variants[0]
    planted.text = planted.text.replace("12%", "45%")
    report = verify_master(master, sources)
    assert not report.ok
    assert any("wording not found" in problem for problem in report.problems)


def test_verify_requires_every_question_answered(sources):
    master, _ = built(sources)
    report = verify_master(master, sources)
    assert any("c_president" in problem for problem in report.problems)


@pytest.mark.parametrize("text, bad", [
    ("Led the team — and delivered", True),
    ("We didn't stop", True),
    ("It's done", True),
    ("Founder's playbook and the CEO's goals", False),
    ("Delivered in 2024 - 2025", False),
])
def test_style_rule(text, bad):
    assert bool(style_problems(text)) is bad


def test_confirmation_is_tied_to_content(tmp_path, sources):
    groups = spec()
    groups.conflicts[0].resolution = "a"
    master, _ = built(sources, groups)
    path = tmp_path / "master_resume.yaml"
    save_master(path, master)
    with pytest.raises(ConfigError, match="not confirmed"):
        require_confirmed(path)

    master.meta.status, master.meta.content_hash = "confirmed", master.content_hash()
    save_master(path, master)
    assert require_confirmed(path).name == "Test Person"

    edited = load_master(path)
    edited.experience[0].title = "Senior Analyst"  # an inflated title slipped in after confirmation
    save_master(path, edited)
    with pytest.raises(ConfigError, match="changed after it was confirmed"):
        require_confirmed(path)


def test_confirmed_resume_never_uses_disputed_wordings(sources):
    master, _ = built(sources)
    profile = load_profile(FIXTURES / "profile.yaml")

    draft = base_document(master, profile)
    assert draft.draft_note and any(b.disputed for s in draft.sections for e in s.entries for b in e.bullets)

    master.meta.status = "confirmed"
    final = base_document(master, profile)
    assert final.draft_note is None
    texts = " ".join(final.all_text())
    assert "President" not in texts
    assert not any(b.disputed for s in final.sections for e in s.entries for b in e.bullets)


def test_file_names():
    assert resume_filename("Test Person", "Spark AI, Inc.") == "Test_Person_Resume_Spark_AI_Inc.pdf"
    assert resume_filename("Test Person") == "Test_Person_Resume.pdf"


def test_style_violation_blocks_rendering(tmp_path, sources):
    master, _ = built(sources)
    master.experience[0].title = "Analyst — Data"
    doc = base_document(master, load_profile(FIXTURES / "profile.yaml"))
    with pytest.raises(RenderError, match="em dash"):
        render_pdf(doc, tmp_path / "x.pdf", tmp_path)


BROWSERS = Path(__file__).resolve().parent.parent / ".playwright-browsers"


@pytest.mark.skipif(not BROWSERS.exists(), reason="Playwright browser not installed")
def test_pdf_is_one_page(tmp_path, sources):
    master, _ = built(sources)
    doc = base_document(master, load_profile(FIXTURES / "profile.yaml"))
    result = render_pdf(doc, tmp_path / resume_filename("Test Person", "Acme"), BROWSERS)
    assert result.path.exists() and result.pages == 1
